#!/usr/bin/env python3
"""
app.py - the XKey website crawler web app.

A visitor enters a website (and optionally a competitor). Crawls run in
background threads; the page shows live progress, then an SEO issues report
with a score out of 100. Results download as Excel, CSV or an issues list,
or open in Google Sheets. Optional extras, each switched on by settings:
weekly/monthly re-crawls that email what changed, and asking for an email
before download.

Run locally:
    pip install -r requirements.txt
    python app.py                       # http://127.0.0.1:5000

Production (one process, many threads: jobs live in memory):
    gunicorn -w 1 --threads 8 -b 0.0.0.0:$PORT app:app

Settings (environment variables) - see README.md for the full list:
    MAX_CONCURRENT_JOBS   crawls allowed at once (default 3)
    JOBS_DIR / DATA_DIR   where crawl files / the SQLite database live
    JOB_RETENTION_HOURS   delete finished crawl files after this (default 48)
    SECRET_KEY            signs visitor sessions (set a long random value)
    PUBLIC_URL            e.g. https://crawler.xkey.co.uk (used in emails)
    SMTP_HOST, SMTP_PORT, SMTP_USER, SMTP_PASSWORD, MAIL_FROM
                          turn on "email me when the site changes"
    SCHEDULE_MAX_PAGES    page limit for scheduled re-crawls (default 1000)
    REQUIRE_EMAIL         "1" asks for an email before downloads (default off)
    ADMIN_PASSWORD        password for /admin/leads.csv (user name: admin)
    ALLOW_PRIVATE_HOSTS   "1" only for local testing
"""

import hmac
import ipaddress
import os
import re
import secrets
import socket
import threading
import time
import uuid
from urllib.parse import urlparse

import requests
from flask import (Flask, Response, abort, jsonify, redirect, render_template, request,
                   send_file, session, url_for)

import mailer
import reports
import store
from seo_crawler import COLUMNS, DEFAULT_MAX_PAGES, HARD_MAX_PAGES, Crawler

MAX_CONCURRENT_JOBS = int(os.environ.get("MAX_CONCURRENT_JOBS", "3"))
MAX_JOBS_PER_VISITOR = int(os.environ.get("MAX_JOBS_PER_VISITOR", "2"))
JOBS_DIR = os.path.abspath(os.environ.get("JOBS_DIR", "jobs"))
JOB_RETENTION_HOURS = float(os.environ.get("JOB_RETENTION_HOURS", "48"))
SCHEDULE_MAX_PAGES = min(int(os.environ.get("SCHEDULE_MAX_PAGES", "1000")), HARD_MAX_PAGES)
REQUIRE_EMAIL = os.environ.get("REQUIRE_EMAIL") == "1"
ADMIN_PASSWORD = os.environ.get("ADMIN_PASSWORD", "")
PUBLIC_URL = os.environ.get("PUBLIC_URL", "").rstrip("/")
ALLOW_PRIVATE_HOSTS = os.environ.get("ALLOW_PRIVATE_HOSTS") == "1"
# Test hook: near-zero delays for local testing only.
FAST_TEST = os.environ.get("CRAWLER_FAST_TEST") == "1"
EMAIL_RE = re.compile(r"^[^@\s]{1,64}@[^@\s]+\.[A-Za-z]{2,}$")

os.makedirs(JOBS_DIR, exist_ok=True)
store.init()
app = Flask(__name__)
# Without SECRET_KEY, sessions reset on every restart (only matters for REQUIRE_EMAIL).
app.secret_key = os.environ.get("SECRET_KEY") or secrets.token_hex(32)

jobs = {}      # job_id -> {"crawler", "created", "url", "ip", "compare_id", "report"}
compares = {}  # compare_id -> {"a": job_id, "b": job_id, "created"}
jobs_lock = threading.Lock()


# ------------------------------------------------------------------ helpers
def pick_scheme(bare):
    """For 'example.com' typed without http(s)://, prefer https, fall back to http."""
    try:
        requests.head("https://" + bare, timeout=10, allow_redirects=True)
        return "https://" + bare
    except requests.RequestException:
        return "http://" + bare


def validate_target(raw):
    """
    Return a clean http(s) URL, or raise ValueError with a friendly message.
    Refuses localhost / private IPs so visitors can't use the server to
    probe its own internal network.
    """
    raw = (raw or "").strip()
    if not raw:
        raise ValueError("Please enter a website address.")
    if "://" not in raw:
        raw = pick_scheme(raw)
    parts = urlparse(raw)
    if parts.scheme not in ("http", "https") or not parts.hostname:
        raise ValueError("That doesn't look like a website address.")
    if not ALLOW_PRIVATE_HOSTS:
        try:
            infos = socket.getaddrinfo(parts.hostname, None)
        except socket.gaierror:
            raise ValueError(f"{parts.hostname} could not be found.")
        for info in infos:
            ip = ipaddress.ip_address(info[4][0])
            if ip.is_private or ip.is_loopback or ip.is_link_local or ip.is_reserved:
                raise ValueError("That address is not allowed.")
    return raw


def visitor_ip():
    # Render (and most hosts) put the real visitor first in X-Forwarded-For.
    forwarded = request.headers.get("X-Forwarded-For", "")
    return forwarded.split(",")[0].strip() or request.remote_addr or "?"


def is_running(job):
    return job["crawler"].state in ("pending", "running")


def host_of(url):
    return urlparse(url).hostname or url


def base_url():
    return PUBLIC_URL or request.url_root.rstrip("/")


def safe_next(value, fallback):
    """Only allow redirects to paths on this site."""
    return value if value and value.startswith("/") and not value.startswith("//") else fallback


def make_crawler(job_id, url, max_pages, path, products_only, out_dir=JOBS_DIR, sitemap_list=False):
    delays = (0.01, 0.02) if FAST_TEST else (2.5, 5.5)
    return Crawler(url, os.path.join(out_dir, f"{job_id}.csv"), max_pages,
                   min_delay=delays[0], max_delay=delays[1],
                   backoff_schedule=[1, 2] if FAST_TEST else None,
                   path_prefix=path or None, products_only=products_only,
                   sitemap_list=sitemap_list)


def start_job(url, max_pages, path, products_only, ip, compare_id=None, sitemap_list=False):
    job_id = uuid.uuid4().hex
    crawler = make_crawler(job_id, url, max_pages, path, products_only, sitemap_list=sitemap_list)
    job = {"crawler": crawler, "created": time.time(), "url": url, "ip": ip,
           "compare_id": compare_id, "report": None, "lock": threading.Lock()}
    jobs[job_id] = job

    def run():
        crawler.run()
        # Build the report and spreadsheets here, not inside a page request: big crawls
        # take a while on a small server and would otherwise freeze the progress page.
        if crawler.pages_done:
            try:
                build_report(job, final=True)
            except Exception as exc:
                crawler.log(f"Could not build the report: {exc}")
    threading.Thread(target=run, daemon=True).start()
    return job_id


def cleanup_old_jobs():
    cutoff = time.time() - JOB_RETENTION_HOURS * 3600
    with jobs_lock:
        for job_id in [k for k, j in jobs.items() if j["created"] < cutoff and not is_running(j)]:
            base = jobs[job_id]["crawler"].output_csv[:-4]
            for ext in (".csv", ".xlsx", "-issues.csv"):
                try:
                    os.remove(base + ext)
                except OSError:
                    pass
            del jobs[job_id]
        for cid in [k for k, c in compares.items() if c["a"] not in jobs or c["b"] not in jobs]:
            try:
                os.remove(os.path.join(JOBS_DIR, f"compare-{cid}.xlsx"))
            except OSError:
                pass
            del compares[cid]


def get_job(job_id):
    job = jobs.get(job_id)
    if not job:
        abort(404)
    return job


def job_files(job):
    base = job["crawler"].output_csv[:-4]
    return {"csv": base + ".csv", "xlsx": base + ".xlsx", "issues": base + "-issues.csv"}


def build_report(job, final):
    """
    Issues, score and spreadsheets for a job. Finished jobs are cached;
    running jobs are rebuilt on each download so partial results work.
    """
    if job["report"] and final:
        return job["report"]
    with job.setdefault("lock", threading.Lock()):
        if job["report"] and final:
            return job["report"]
        return _build_report(job, final)


def _build_report(job, final):
    files = job_files(job)
    rows = reports.load_rows(files["csv"])
    issues = reports.find_issues(rows)
    summary = reports.summarise(rows, issues)
    reports.write_issues_csv(issues, files["issues"])
    reports.write_xlsx(rows, files["xlsx"], issues, summary, site=job["url"])
    report = {"summary": summary, "top_issues": issues[:15], "total_issues": len(issues)}
    if final:
        job["report"] = report
    return report


def unlocked(job_id):
    return not REQUIRE_EMAIL or job_id in session.get("unlocked", [])


# ------------------------------------------------------------------- pages
@app.get("/healthz")
def healthz():
    """Used by the host (e.g. Render) to check the app is up."""
    return "ok"


def render_form(form=None, error=None, code=200):
    return render_template("index.html", form=form or {}, error=error,
                           default_max=DEFAULT_MAX_PAGES, hard_max=HARD_MAX_PAGES,
                           email_on=mailer.configured(), schedule_max=SCHEDULE_MAX_PAGES), code


@app.get("/")
def index():
    return render_form()


@app.post("/start")
def start():
    cleanup_old_jobs()
    form = {k: (request.form.get(k) or "").strip() for k in
            ("url", "max_pages", "path", "products_only", "product_list", "competitor", "watch_email", "every_days")}
    try:
        max_pages = max(1, min(int(form["max_pages"] or DEFAULT_MAX_PAGES), HARD_MAX_PAGES))
    except ValueError:
        max_pages = DEFAULT_MAX_PAGES
    path, products_only = form["path"][:200], form["products_only"] == "1"
    product_list = form["product_list"] == "1"
    if product_list:
        products_only = True

    try:
        target = validate_target(form["url"])
        competitor = validate_target(form["competitor"]) if form["competitor"] else None
    except ValueError as exc:
        return render_form(form, str(exc), 400)

    watch_email = form["watch_email"].lower()
    if watch_email and mailer.configured():
        if not EMAIL_RE.match(watch_email):
            return render_form(form, "Please enter a valid email address.", 400)
        if store.active_count_for(watch_email) >= 3:
            return render_form(form, "That email already has 3 sites being watched.", 400)

    ip = visitor_ip()
    needed = 2 if competitor else 1
    with jobs_lock:
        running = [j for j in jobs.values() if is_running(j)]
        if len(running) + needed > MAX_CONCURRENT_JOBS:
            return render_form(form, "The crawler is busy right now. Please try again in a few minutes.", 503)
        if sum(1 for j in running if j["ip"] == ip) + needed > MAX_JOBS_PER_VISITOR:
            return render_form(form, "You already have crawls running. Please wait for them to finish.", 429)
        compare_id = uuid.uuid4().hex if competitor else None
        job_a = start_job(target, max_pages, path, products_only, ip, compare_id, sitemap_list=product_list)
        if competitor:
            job_b = start_job(competitor, max_pages, path, products_only, ip, compare_id, sitemap_list=product_list)
            compares[compare_id] = {"a": job_a, "b": job_b, "created": time.time()}

    if watch_email and mailer.configured():
        every = 30 if form["every_days"] == "30" else 7
        token = store.add_schedule(watch_email, target, min(max_pages, SCHEDULE_MAX_PAGES),
                                   path, products_only, every)
        link = f"{base_url()}{url_for('schedule_confirm', token=token)}"
        mailer.send(watch_email, f"Confirm: watch {host_of(target)} for changes",
                    f"Someone (hopefully you) asked XKey to re-crawl {target} every "
                    f"{'month' if every == 30 else 'week'} and email you what changed.\n\n"
                    f"Confirm here: {link}\n\nIf this wasn't you, ignore this email "
                    f"and nothing will be sent.")

    if competitor:
        return redirect(url_for("compare_page", compare_id=compare_id))
    return redirect(url_for("job_page", job_id=job_a))


@app.get("/job/<job_id>")
def job_page(job_id):
    job = get_job(job_id)
    return render_template("job.html", job_id=job_id, host=host_of(job["url"]),
                           locked=not unlocked(job_id), compare_id=job["compare_id"],
                           retention=int(JOB_RETENTION_HOURS),
                           sheet_url=f"{base_url()}{url_for('sheet_csv', job_id=job_id)}")


@app.get("/api/job/<job_id>")
def job_status(job_id):
    job = get_job(job_id)
    c = job["crawler"]
    report = job["report"]
    state = c.state
    if not is_running(job) and c.pages_done and report is None and job.get("lock") and job["lock"].locked():
        state = "preparing"  # crawl done, spreadsheets still being built in the background
    return jsonify(state=state, message=c.message, pages_done=c.pages_done,
                   mode="list" if c.sitemap_list else "crawl",
                   pages_fetched=c.pages_fetched, max_pages=c.max_pages, queued=c.queued,
                   log=list(c.recent_log)[-15:], report=report)


@app.post("/job/<job_id>/stop")
def stop(job_id):
    get_job(job_id)["crawler"].stop()
    return redirect(url_for("job_page", job_id=job_id))


@app.get("/job/<job_id>/download/<kind>")
def download(job_id, kind):
    job = get_job(job_id)
    if kind not in ("csv", "xlsx", "issues"):
        abort(404)
    if not unlocked(job_id):
        return redirect(url_for("unlock", job_id=job_id))
    files = job_files(job)
    if not os.path.exists(files["csv"]):
        with open(files["csv"], "w", encoding="utf-8-sig") as f:
            f.write(",".join(COLUMNS) + "\n")
    if kind != "csv":
        build_report(job, final=not is_running(job))
    host = host_of(job["url"])
    names = {"csv": f"crawl-{host}.csv", "xlsx": f"crawl-{host}.xlsx", "issues": f"seo-issues-{host}.csv"}
    mime = ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            if kind == "xlsx" else "text/csv")
    return send_file(files[kind], mimetype=mime, as_attachment=True, download_name=names[kind])


@app.get("/job/<job_id>/sheet.csv")
def sheet_csv(job_id):
    """Plain CSV for Google Sheets =IMPORTDATA(). The job id is the secret."""
    path = job_files(get_job(job_id))["csv"]
    if not os.path.exists(path):
        abort(404)
    return send_file(path, mimetype="text/csv")


@app.route("/job/<job_id>/unlock", methods=["GET", "POST"])
def unlock(job_id):
    job = get_job(job_id)
    next_url = safe_next(request.values.get("next"), url_for("job_page", job_id=job_id))
    if not REQUIRE_EMAIL:
        return redirect(next_url)
    error = None
    if request.method == "POST":
        email = (request.form.get("email") or "").strip().lower()
        if EMAIL_RE.match(email):
            store.add_lead(email, job["url"], job_id, request.form.get("marketing_ok") == "1")
            ids = set(session.get("unlocked", []))
            ids.add(job_id)
            cid = job["compare_id"]
            if cid and cid in compares:
                ids.update((compares[cid]["a"], compares[cid]["b"]))
            session["unlocked"] = list(ids)[-50:]
            return redirect(next_url)
        error = "Please enter a valid email address."
    return render_template("unlock.html", host=host_of(job["url"]), next_url=next_url, error=error)


# ----------------------------------------------------------------- compare
def get_compare(compare_id):
    cmp_ = compares.get(compare_id)
    if not cmp_ or cmp_["a"] not in jobs or cmp_["b"] not in jobs:
        abort(404)
    return cmp_, jobs[cmp_["a"]], jobs[cmp_["b"]]


@app.get("/compare/<compare_id>")
def compare_page(compare_id):
    cmp_, a, b = get_compare(compare_id)
    return render_template("compare.html", compare_id=compare_id, job_a=cmp_["a"], job_b=cmp_["b"],
                           host_a=host_of(a["url"]), host_b=host_of(b["url"]),
                           locked=not unlocked(cmp_["a"]))


@app.get("/api/compare/<compare_id>")
def compare_status(compare_id):
    _, a, b = get_compare(compare_id)

    def brief(job):
        c = job["crawler"]
        return {"state": c.state, "pages_done": c.pages_done, "max_pages": c.max_pages}

    done = not is_running(a) and not is_running(b)
    rows = None
    if done:
        ra = build_report(a, final=True) if a["crawler"].pages_done else {"summary": {}}
        rb = build_report(b, final=True) if b["crawler"].pages_done else {"summary": {}}
        rows = reports.compare(ra["summary"], rb["summary"])
    return jsonify(a=brief(a), b=brief(b), done=done, rows=rows,
                   host_a=host_of(a["url"]), host_b=host_of(b["url"]))


@app.get("/compare/<compare_id>/download")
def compare_download(compare_id):
    cmp_, a, b = get_compare(compare_id)
    if not unlocked(cmp_["a"]):
        return redirect(url_for("unlock", job_id=cmp_["a"],
                                next=url_for("compare_page", compare_id=compare_id)))
    from openpyxl import Workbook
    from openpyxl.styles import Font
    sa = build_report(a, final=not is_running(a))["summary"]
    sb = build_report(b, final=not is_running(b))["summary"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Comparison"
    ws.append(["", host_of(a["url"]), host_of(b["url"]), "Better"])
    for metric, va, vb, winner in reports.compare(sa, sb):
        ws.append([metric, va, vb, {"a": host_of(a["url"]), "b": host_of(b["url"])}.get(winner, "")])
    ws.append([])
    ws.append(["Issue", host_of(a["url"]), host_of(b["url"])])
    for issue in sorted(set(sa["issue_counts"]) | set(sb["issue_counts"])):
        ws.append([issue, sa["issue_counts"].get(issue, 0), sb["issue_counts"].get(issue, 0)])
    for row in ws.iter_rows():
        row[0].font = Font(bold=True)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for col, width in zip("ABCD", (34, 26, 26, 26)):
        ws.column_dimensions[col].width = width
    path = os.path.join(JOBS_DIR, f"compare-{compare_id}.xlsx")
    wb.save(path)
    return send_file(path, as_attachment=True,
                     download_name=f"compare-{host_of(a['url'])}-vs-{host_of(b['url'])}.xlsx")


# --------------------------------------------------------------- schedules
@app.get("/watch/confirm/<token>")
def schedule_confirm(token):
    sched = store.get_schedule(token)
    if not sched:
        abort(404)
    store.confirm_schedule(token)
    return render_template("message.html", heading="You're all set",
                           text=f"We'll re-crawl {sched['url']} every "
                                f"{'month' if sched['every_days'] == 30 else 'week'} and email "
                                f"{sched['email']} what changed. The first report is on its way.")


@app.get("/watch/stop/<token>")
def schedule_stop(token):
    sched = store.get_schedule(token)
    if not sched:
        abort(404)
    store.stop_schedule(token)
    return render_template("message.html", heading="Stopped",
                           text=f"We won't email you about {sched['url']} again.")


def run_due_schedules():
    """Re-crawl each due site, email what changed, keep the new crawl."""
    store.delete_unconfirmed()
    sched_dir = os.path.join(store.DATA_DIR, "schedules")
    os.makedirs(sched_dir, exist_ok=True)
    for sched in store.due_schedules():
        name = f"{sched['id']}-{int(time.time())}"
        crawler = make_crawler(name, sched["url"], sched["max_pages"], sched["path"],
                               bool(sched["products_only"]), out_dir=sched_dir)
        crawler.run()
        if not crawler.pages_done:
            # Site unreachable this time: try again in a day rather than losing a week.
            store.finish_run(sched["id"], sched["last_csv"], 1)
            continue
        new_rows = reports.load_rows(crawler.output_csv)
        old_rows = reports.load_rows(sched["last_csv"]) if sched["last_csv"] else []
        summary = reports.summarise(new_rows)
        xlsx = crawler.output_csv[:-4] + ".xlsx"
        reports.write_xlsx(new_rows, xlsx, summary=summary, site=sched["url"])
        changes = reports.diff_crawls(old_rows, new_rows) if old_rows else None
        stop_link = f"{PUBLIC_URL}/watch/stop/{sched['token']}" if PUBLIC_URL else ""
        body = (f"XKey re-crawled {sched['url']} ({len(new_rows)} pages). "
                + (f"SEO score: {summary['Score']}/100.\n\n" if summary["Score"] != reports.MISSING
                   else "No SEO score: no pages could be read.\n\n")
                + ("This is the first crawl, so there's nothing to compare yet. "
                   "Next time we'll list what changed." if changes is None
                   else reports.diff_as_text(changes))
                + "\n\nThe full spreadsheet is attached."
                + (f"\n\nStop these emails: {stop_link}" if stop_link else ""))
        mailer.send(sched["email"], f"{host_of(sched['url'])}: "
                    + ("first crawl report" if changes is None
                       else f"{sum(len(v) for v in changes.values())} changes"), body, xlsx)
        # Keep only the latest crawl per schedule.
        if sched["last_csv"] and os.path.exists(sched["last_csv"]):
            os.remove(sched["last_csv"])
        try:
            os.remove(xlsx)
        except OSError:
            pass
        store.finish_run(sched["id"], crawler.output_csv, sched["every_days"])


def scheduler_loop():
    while True:
        try:
            run_due_schedules()
        except Exception as exc:  # keep the scheduler alive whatever happens
            print(f"Scheduler error: {exc}", flush=True)
        time.sleep(2 if FAST_TEST else 300)


if mailer.configured() and os.environ.get("DISABLE_SCHEDULER") != "1":
    threading.Thread(target=scheduler_loop, daemon=True).start()


# ------------------------------------------------------------------- admin
@app.get("/admin/leads.csv")
def admin_leads():
    auth = request.authorization
    if not ADMIN_PASSWORD or not auth or auth.username != "admin" \
            or not hmac.compare_digest(auth.password or "", ADMIN_PASSWORD):
        return Response("Login required", 401, {"WWW-Authenticate": 'Basic realm="admin"'})
    return Response(store.leads_csv(), mimetype="text/csv",
                    headers={"Content-Disposition": "attachment; filename=leads.csv"})


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=int(os.environ.get("PORT", "5000")), threaded=True)
